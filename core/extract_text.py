# PDF / DOCX / TXT -> plain string. Never crashes on a bad page.

import pathlib
import re

from core.storage import load_company_data

ROLE_TITLES = [
    "Associate Software Engineer",
    "Application Support Analyst",
    "Software Engineer",
    "Data Scientist",
    "Data Analyst",
]


def extract_text(path: str) -> str:
    suffix = pathlib.Path(path).suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix in (".txt", ".md"):
        return pathlib.Path(path).read_text(encoding="utf-8", errors="replace")

    raise ValueError(f"Unsupported file type {suffix!r} for {path} — use PDF, DOCX, or TXT")


def _extract_pdf(path: str) -> str:
    import pypdf

    reader = pypdf.PdfReader(path)
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(pages)


def _extract_docx(path: str) -> str:
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def guess_company_and_role(text: str, filename: str) -> tuple[str, str]:
    window = text[:2000]
    window_lower = re.sub(r"\s+", " ", window.lower())

    company = ""
    companies = load_company_data()
    for name in companies:
        if name.lower() in window_lower:
            company = name
            break

    role = ""
    for title in ROLE_TITLES:
        if title.lower() in window_lower:
            role = title
            break

    if not company or not role:
        stem = pathlib.Path(filename).stem
        if "-" in stem:
            parts = re.split(r"\s*-\s*", stem, maxsplit=1)
        else:
            stem_clean = re.sub(r"[_\-]+", " ", stem).strip()
            parts = stem_clean.split(" ", 1)
        if not company and parts:
            company = parts[0].strip()
        if not role and len(parts) > 1:
            role = parts[1].strip()

    return company, role
